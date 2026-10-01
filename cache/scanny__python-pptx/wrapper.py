import os
import uuid
from typing import List, Dict, Any

from mcp.server.fastmcp import FastMCP

mcp = FastMCP("python-pptx-tools")


def _abs(path: str) -> str:
    return os.path.abspath(path)


@mcp.tool()
def create_presentation(title: str, subtitle: str = "", output_path: str | None = None) -> str:
    """
    Create a new PPTX presentation with a title slide.

    Args:
        title: Title text for the first slide.
        subtitle: Optional subtitle text.
        output_path: Optional path to save the PPTX. If not provided, a file named generated_<uuid>.pptx is created.

    Returns:
        Absolute path to the created PPTX file.
    """
    from pptx import Presentation

    if output_path is None:
        output_path = f"./generated_{uuid.uuid4().hex}.pptx"

    prs = Presentation()
    # Title Slide layout is usually index 0
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)

    # Set title and subtitle if placeholders are present
    if hasattr(slide.shapes, "title") and slide.shapes.title is not None:
        slide.shapes.title.text = title
    # Subtitle is typically placeholder index 1
    try:
        subtitle_placeholder = slide.placeholders[1]
        subtitle_placeholder.text = subtitle or ""
    except Exception:
        # If layout doesn't have a subtitle placeholder, ignore
        pass

    prs.save(output_path)
    return _abs(output_path)


@mcp.tool()
def add_text_slide(pptx_path: str, title: str, bullets: List[str] | None = None, output_path: str | None = None) -> str:
    """
    Add a 'Title and Content' text slide with optional bullet points to an existing PPTX.
    Writes to output_path or a new *_modified.pptx to avoid overwriting read-only inputs.

    Args:
        pptx_path: Path to the source PPTX file (read-only).
        title: Title for the new slide.
        bullets: Optional list of bullet point strings.
        output_path: Optional path for the modified PPTX. Defaults to <pptx_path stem>_modified.pptx.

    Returns:
        Absolute path to the saved modified PPTX file.
    """
    from pptx import Presentation

    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"PPTX not found: {pptx_path}")

    if output_path is None:
        base, ext = os.path.splitext(pptx_path)
        output_path = f"{base}_modified{ext or '.pptx'}"

    prs = Presentation(pptx_path)

    # 'Title and Content' layout is commonly index 1
    layout_index = 1 if len(prs.slide_layouts) > 1 else 0
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])

    # Set title if available
    if hasattr(slide.shapes, "title") and slide.shapes.title is not None:
        slide.shapes.title.text = title

    # Add bullets into the main content placeholder if present
    try:
        body = slide.placeholders[1]
        tf = body.text_frame
        if bullets and len(bullets) > 0:
            tf.text = bullets[0]
            for b in bullets[1:]:
                p = tf.add_paragraph()
                p.text = b
                p.level = 0
        else:
            tf.text = ""
    except Exception:
        # If no suitable placeholder, ignore bullets
        pass

    prs.save(output_path)
    return _abs(output_path)


@mcp.tool()
def extract_text(pptx_path: str) -> List[Dict[str, Any]]:
    """
    Extract text content from all slides in a PPTX.

    Args:
        pptx_path: Path to the PPTX file.

    Returns:
        A list of dicts, each with:
          - slide_index: 1-based slide number
          - texts: list of text snippets found on the slide
    """
    from pptx import Presentation

    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"PPTX not found: {pptx_path}")

    prs = Presentation(pptx_path)
    results: List[Dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts: List[str] = []
        # Collect text from shapes with text frames
        for shape in slide.shapes:
            try:
                if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                    text = shape.text
                    if text:
                        # Split lines to keep content granular
                        for line in text.splitlines():
                            if line.strip():
                                texts.append(line.strip())
                # Also collect from tables if present
                if hasattr(shape, "has_table") and shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        for cell in row.cells:
                            for line in cell.text.splitlines():
                                if line.strip():
                                    texts.append(line.strip())
            except Exception:
                # Skip shapes that can't be interrogated safely
                continue

        results.append({"slide_index": idx, "texts": texts})

    return results


@mcp.tool()
def get_presentation_info(pptx_path: str) -> Dict[str, Any]:
    """
    Get basic info about a PPTX: slide count and slide titles (if any).

    Args:
        pptx_path: Path to the PPTX file.

    Returns:
        Dict with keys:
          - slide_count: number of slides
          - titles: list of title texts (empty string if none on a slide)
    """
    from pptx import Presentation

    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"PPTX not found: {pptx_path}")

    prs = Presentation(pptx_path)
    titles: List[str] = []
    for slide in prs.slides:
        t = ""
        try:
            if hasattr(slide.shapes, "title") and slide.shapes.title is not None:
                t = slide.shapes.title.text or ""
        except Exception:
            t = ""
        titles.append(t)
    return {"slide_count": len(prs.slides), "titles": titles}


if __name__ == "__main__":
    # Self-test (offline): create a PPTX, verify it's written, and extract some text.
    path = create_presentation("Hello from python-pptx", "This is a subtitle")
    if not os.path.exists(path) or os.path.getsize(path) <= 0:
        raise RuntimeError("Self-test failed: created PPTX not found or empty.")

    extracted = extract_text(path)
    if not extracted:  # Must be non-empty result
        raise RuntimeError("Self-test failed: extract_text returned empty result.")

    # Also test adding a text slide
    mod_path = add_text_slide(path, "Agenda", ["First", "Second"])
    if not os.path.exists(mod_path) or os.path.getsize(mod_path) <= 0:
        raise RuntimeError("Self-test failed: modified PPTX not found or empty.")

    info = get_presentation_info(mod_path)
    if not info or "slide_count" not in info:
        raise RuntimeError("Self-test failed: get_presentation_info returned empty or invalid result.")

    print("Self-test OK")