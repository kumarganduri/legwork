import os
import time
from typing import List, Optional, Dict, Any

from mcp.server.fastmcp import FastMCP
from docx import Document

mcp = FastMCP("python-docx")


def _ensure_docx_path(out_path: Optional[str], prefix: str = "document") -> str:
    """
    Ensure output path is a .docx under the current working directory.
    If out_path is None, generate a unique filename.
    If out_path is provided, only use its basename to keep output under CWD.
    """
    cwd = os.getcwd()
    if out_path:
        base = os.path.basename(out_path)
        if not base.lower().endswith(".docx"):
            base += ".docx"
    else:
        base = f"{prefix}-{int(time.time()*1000)}.docx"
    return os.path.abspath(os.path.join(cwd, base))


@mcp.tool()
def create_docx(paragraphs: List[str], out_path: Optional[str] = None, heading: Optional[str] = None) -> Dict[str, Any]:
    """
    Create a new .docx document with the given paragraphs (and optional heading) and save it.

    Args:
        paragraphs: List of paragraph texts to add in order.
        out_path: Optional desired filename. Will be placed under the current directory.
        heading: Optional heading text to add at the top (Heading 1 style).

    Returns:
        {"path": "<absolute path to saved .docx>"}
    """
    save_path = _ensure_docx_path(out_path, prefix="new")
    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs:
        doc.add_paragraph("" if p is None else str(p))
    doc.save(save_path)
    return {"path": os.path.abspath(save_path)}


@mcp.tool()
def read_docx_paragraphs(path: str) -> Dict[str, Any]:
    """
    Read all paragraph texts from a .docx file.

    Args:
        path: Path to an existing .docx file.

    Returns:
        {"paragraphs": ["para1", "para2", ...]}
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(f"No such file: {path}")
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs]
    return {"paragraphs": paras}


@mcp.tool()
def append_paragraph_copy(src_path: str, text: str, out_path: Optional[str] = None, style: Optional[str] = None) -> Dict[str, Any]:
    """
    Append a paragraph to a copy of an existing .docx, writing the new file under the current directory.

    Args:
        src_path: Path to the source .docx (read-only).
        text: Paragraph text to append.
        out_path: Optional new filename for the copied/modified document.
        style: Optional paragraph style name (e.g., 'Normal', 'Heading 2').

    Returns:
        {"path": "<absolute path to new .docx>", "paragraph_count": <int>}
    """
    if not os.path.isfile(src_path):
        raise FileNotFoundError(f"No such file: {src_path}")
    save_path = _ensure_docx_path(out_path, prefix="appended")
    doc = Document(src_path)
    para = doc.add_paragraph("" if text is None else str(text))
    if style:
        try:
            para.style = style
        except Exception:
            # If the style doesn't exist, leave default; don't fail the entire operation.
            pass
    doc.save(save_path)
    # Reopen to count reliably
    new_doc = Document(save_path)
    return {"path": os.path.abspath(save_path), "paragraph_count": len(new_doc.paragraphs)}


if __name__ == "__main__":
    # Self-test: create a small document and read it back.
    try:
        created = create_docx(paragraphs=["Hello, python-docx MCP!"], heading="Self Test")
        created_path = created.get("path")
        if not created_path or not os.path.isfile(created_path) or os.path.getsize(created_path) <= 0:
            raise RuntimeError("Self-test: created file missing or empty")

        readback = read_docx_paragraphs(created_path)
        paras = readback.get("paragraphs", [])
        if not paras or not any(p.strip() for p in paras):
            raise RuntimeError("Self-test: no paragraphs read back")

        # Also test append to copy
        appended = append_paragraph_copy(created_path, "Appended line.")
        appended_path = appended.get("path")
        if not appended_path or not os.path.isfile(appended_path) or os.path.getsize(appended_path) <= 0:
            raise RuntimeError("Self-test: appended file missing or empty")

        print("OK")
    except Exception as e:
        raise SystemExit(f"Self-test failed: {e}")